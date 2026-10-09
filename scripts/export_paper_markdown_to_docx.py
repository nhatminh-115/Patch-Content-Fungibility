"""Export the PCF manuscript to an editable IEEE-numbered Word document."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FENCE = chr(96) * 3


def bibliography(markdown: str):
    block = re.search(re.escape(FENCE) + r"bibtex\s*\n(.*?)\n" + re.escape(FENCE), markdown, re.S)
    if not block:
        raise ValueError("BibTeX block is missing")
    source = block.group(1)
    entries, pos = [], 0
    while True:
        start = source.find("@", pos)
        if start < 0:
            break
        brace = source.find("{", start)
        if brace < 0:
            break
        depth, end = 0, None
        for i in range(brace, len(source)):
            if source[i] == "{":
                depth += 1
            elif source[i] == "}":
                depth -= 1
                if depth == 0:
                    end = i + 1
                    break
        if end is None:
            raise ValueError("Unbalanced BibTeX entry")
        comma = source.find(",", brace)
        head = re.match(r"@(\w+)\s*\{\s*([^,]+),", source[start:comma + 1]) if comma >= 0 else None
        if not head:
            raise ValueError("Could not parse BibTeX entry")
        key, body = head.group(2).strip(), source[brace + 1:end - 1]
        fields, i = {}, 0
        while i < len(body):
            match = re.search(r"([A-Za-z]+)\s*=\s*", body[i:])
            if not match:
                break
            name, j = match.group(1).lower(), i + match.end()
            while j < len(body) and body[j].isspace():
                j += 1
            if j >= len(body):
                break
            if body[j] == "{":
                d, k = 1, j + 1
                while k < len(body) and d:
                    if body[k] == "{":
                        d += 1
                    elif body[k] == "}":
                        d -= 1
                    k += 1
                value = body[j + 1:k - 1]
            elif body[j] == '"':
                k = j + 1
                while k < len(body) and not (body[k] == '"' and body[k - 1] != "\\"):
                    k += 1
                value = body[j + 1:k]
            else:
                k = j
                while k < len(body) and body[k] not in ",\n":
                    k += 1
                value = body[j:k]
            fields[name] = value.strip()
            i = max(k, j) + 1
        entries.append((key, fields))
        pos = end
    if not entries:
        raise ValueError("No BibTeX entries parsed")
    return block, entries


def clean_tex(value: str) -> str:
    replacements = {
        r"\&": "&", r"\_": "_", r"\%": "%", r"\#": "#",
        r"{\'e}": "é", r"{\'E}": "É", r"{\'o}": "ó",
        r"{\'a}": "á", r"{\v{s}}": "š", r"{\v{c}}": "č",
        r"{\"o}": "ö", r"{\"u}": "ü", r"{\~n}": "ñ",
    }
    for old, new in replacements.items():
        value = value.replace(old, new)
    value = re.sub(r"\\textit\{([^{}]*)\}", r"\1", value)
    value = value.replace("{", "").replace("}", "").replace("~", " ")
    return re.sub(r"\s+", " ", value).strip()


def author_names(value: str) -> list[str]:
    return [clean_tex(x.strip()) for x in re.split(r"\s+and\s+", value) if x.strip()]


def initials(name: str) -> str:
    if "," in name:
        family, given = [x.strip() for x in name.split(",", 1)]
    else:
        parts = name.split()
        family, given = parts[-1], " ".join(parts[:-1])
    short = []
    for token in given.split():
        token = token.strip(".")
        if not token:
            continue
        if "-" in token:
            short.append("-".join(x[:1].upper() + "." for x in token.split("-") if x))
        elif len(token) > 1 and token.isupper():
            short.extend(x + "." for x in token)
        else:
            short.append(token[0].upper() + ".")
    return ((" ".join(short) + " ") if short else "") + family


def format_authors(raw: str) -> str:
    names = author_names(raw)
    if len(names) > 6:
        return ", ".join(initials(n) for n in names[:6]) + ", et al."
    if not names:
        return "Author not listed"
    return ", ".join(initials(n) for n in names)


def venue_name(fields: dict[str, str]) -> str:
    venue = clean_tex(fields.get("journal", fields.get("booktitle", "")))
    rules = [
        ("Proceedings of the IEEE/CVF International Conference on Computer Vision", "Proc. IEEE/CVF Int. Conf. Comput. Vis."),
        ("Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition", "Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit."),
        ("Proceedings of the IEEE/CVF Winter Conference on Applications of Computer Vision", "Proc. IEEE/CVF Winter Conf. Appl. Comput. Vis."),
        ("Proceedings of the 34th International Conference on Machine Learning", "Proc. 34th Int. Conf. Mach. Learn."),
        ("Proceedings of the 38th International Conference on Machine Learning", "Proc. 38th Int. Conf. Mach. Learn."),
        ("Proceedings of the 39th International Conference on Machine Learning", "Proc. 39th Int. Conf. Mach. Learn."),
        ("Proceedings of the 40th International Conference on Machine Learning", "Proc. 40th Int. Conf. Mach. Learn."),
        ("International Conference on Learning Representations", "Proc. Int. Conf. Learn. Represent."),
        ("Advances in Neural Information Processing Systems", "Advances in Neural Information Processing Systems"),
        ("Proceedings of the AAAI Conference on Artificial Intelligence", "Proc. AAAI Conf. Artif. Intell."),
        ("International Conference on Neural Networks", "Proc. IEEE Int. Conf. Neural Netw."),
        ("Computer Vision -- ECCV 2022", "Computer Vision – ECCV 2022"),
    ]
    for source, target in rules:
        if source in venue:
            venue = venue.replace(source, target)
            break
    return venue


def ieee_reference(fields: dict[str, str]) -> str:
    authors = format_authors(fields.get("author", ""))
    title = clean_tex(fields.get("title", ""))
    venue = venue_name(fields)
    year = clean_tex(fields.get("year", ""))
    volume = clean_tex(fields.get("volume", ""))
    number = clean_tex(fields.get("number", ""))
    pages = clean_tex(fields.get("pages", "")).replace("--", "–")
    doi = clean_tex(fields.get("doi", ""))
    url = clean_tex(fields.get("url", ""))
    result = f'{authors}, “{title},”'
    if venue:
        result += f" in {venue},"
    if volume:
        result += f" vol. {volume},"
    if number:
        result += f" no. {number},"
    if pages:
        result += f" pp. {pages},"
    result += f" {year}."
    if doi:
        result += f" DOI: {doi}."
    elif url:
        result += f" Available: {url}."
    return result


def cited_order(body: str, entries):
    valid = {key for key, _ in entries}
    order, unresolved = [], []
    for match in re.finditer(r"\\citep\{([^}]+)\}", body):
        for key in (part.strip() for part in match.group(1).split(",")):
            if key not in valid:
                unresolved.append(key)
            elif key not in order:
                order.append(key)
    uncited = sorted(valid - set(order))
    if unresolved or uncited:
        raise ValueError(f"Citation audit failed; unresolved={unresolved}, uncited={uncited}")
    return order


def citation_replace(text: str, key_to_num: dict[str, int]) -> str:
    def cite(match):
        keys = [x.strip() for x in match.group(1).split(",")]
        nums = [key_to_num[k] for k in keys]
        if len(nums) >= 3 and nums == list(range(nums[0], nums[0] + len(nums))):
            return f"[{nums[0]}]–[{nums[-1]}]"
        return ", ".join(f"[{n}]" for n in nums)
    return re.sub(r"\\citep\{([^}]+)\}", cite, text)


def set_font(run, size=10.3, bold=None, italic=None, color="202A32", name="Times New Roman"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


BT = chr(96)
INLINE = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|" + re.escape(BT) + r"(.+?)" + re.escape(BT) + r"|\\citep\{([^}]+)\}|<sub>([^<]+)</sub>|<sup>([^<]+)</sup>|\\\((.+?)\\\))")
INLINE_SCRIPT = re.compile(
    r"(?P<base>\|\|.*?\|\||\([^()]*\)|\{[^{}]*\}|(?:[^\W_]|[\u0300-\u036f])+)"
    r"(?:(?P<sub>_\{[^{}]*\}|_[^\W_])(?P<sup>\^\{[^{}]*\}|\^[^\W_])?|"
    r"(?P<sup_only>\^\{[^{}]*\}|\^[^\W_]))"
)


def normalize_script_glyphs(value: str) -> str:
    """Convert Unicode script glyphs into explicit notation for Word math parsing."""
    super_map = {"⁰": "0", "¹": "1", "²": "2", "³": "3", "⁴": "4",
                 "⁵": "5", "⁶": "6", "⁷": "7", "⁸": "8", "⁹": "9",
                 "⁻": "-", "ᵀ": "T"}
    sub_map = {"₀": "0", "₁": "1", "₂": "2", "₃": "3", "₄": "4",
               "₅": "5", "₆": "6", "₇": "7", "₈": "8", "₉": "9"}
    out = []
    depth = 0
    i = 0
    while i < len(value):
        char = value[i]
        if char == "{":
            depth += 1
            out.append(char)
            i += 1
            continue
        if char == "}":
            depth = max(0, depth - 1)
            out.append(char)
            i += 1
            continue
        if char in super_map:
            chars = []
            while i < len(value) and value[i] in super_map:
                chars.append(super_map[value[i]])
                i += 1
            content = "".join(chars)
            out.append(content if depth else "^{" + content + "}")
            continue
        if char in sub_map:
            chars = []
            while i < len(value) and value[i] in sub_map:
                chars.append(sub_map[value[i]])
                i += 1
            content = "".join(chars)
            out.append(content if depth else "_{" + content + "}")
            continue
        out.append(char)
        i += 1
    return "".join(out)


def add_styled_fragment(p, text, key_to_num, *, size=10.3, bold=None, italic=None, color="202A32"):
    """Append prose and inline symbols, turning every _ and ^ index into Word math."""
    text = re.sub(r"<sub>(.*?)</sub>", r"_{\1}", text)
    text = re.sub(r"<sup>(.*?)</sup>", r"^{\1}", text)
    text = normalize_script_glyphs(text)
    cursor = 0
    for match in INLINE_SCRIPT.finditer(text):
        if match.start() > cursor:
            run = p.add_run(citation_replace(text[cursor:match.start()], key_to_num))
            set_font(run, size=size, bold=bold, italic=italic, color=color)
        math_obj = OxmlElement("m:oMath")
        append_math_expression(math_obj, match.group(0))
        p._p.append(math_obj)
        cursor = match.end()
    if cursor < len(text):
        run = p.add_run(citation_replace(text[cursor:], key_to_num))
        set_font(run, size=size, bold=bold, italic=italic, color=color)


def add_inline_paragraph(doc, text, key_to_num, *, caption=False, center=False, size=10.3):
    p = doc.add_paragraph()
    if caption or center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            add_styled_fragment(p, text[cursor:match.start()], key_to_num,
                                size=8.8 if caption else size, italic=caption,
                                color="4C5963" if caption else "202A32")
        if match.group(2) is not None:
            add_styled_fragment(p, match.group(2), key_to_num, size=size, bold=True)
        elif match.group(3) is not None:
            add_styled_fragment(p, match.group(3), key_to_num,
                                size=8.8 if caption else size, italic=True,
                                color="4C5963" if caption else "202A32")
        elif match.group(4) is not None:
            run = p.add_run(match.group(4))
            set_font(run, size=9.0, name="Consolas", color="34495E")
        elif match.group(6) is not None:
            run = p.add_run(match.group(6))
            set_font(run, size=8.8 if caption else size, italic=caption,
                     color="4C5963" if caption else "202A32")
            run.font.subscript = True
        elif match.group(7) is not None:
            run = p.add_run(match.group(7))
            set_font(run, size=8.8 if caption else size, italic=caption,
                     color="4C5963" if caption else "202A32")
            run.font.superscript = True
        elif match.group(8) is not None:
            math_obj = OxmlElement("m:oMath")
            append_math_expression(math_obj, math_text(match.group(8)))
            p._p.append(math_obj)
        else:
            run = p.add_run(citation_replace(match.group(0), key_to_num))
            set_font(run, size=size, color="202A32")
        cursor = match.end()
    if cursor < len(text):
        add_styled_fragment(p, text[cursor:], key_to_num,
                            size=8.8 if caption else size, italic=caption,
                            color="4C5963" if caption else "202A32")
    p.paragraph_format.space_after = Pt(4 if caption else 5)
    if caption:
        p.paragraph_format.keep_together = True
    if not caption:
        p.paragraph_format.line_spacing = 1.08
    return p


def math_text(source: str) -> str:
    """Normalize equation source while preserving sub/superscript structure."""
    value = re.sub(r"\\tag\{\d+\}", "", source.strip())
    value = normalize_script_glyphs(value)
    value = value.replace("p\u0303", "p~")
    value = value.replace("·", " · ")
    value = value.replace("ℓ", "l").replace("ℝ", "R").replace("∇", "grad")
    for command, replacement in (
        (r"\mathbb{R}", "R"), (r"\operatorname{vec}", "vec"),
        (r"\operatorname{flat}", "flat"), (r"\operatorname{tr}", "tr"),
        (r"\operatorname{dim}", "dim"), (r"\mathrm{mean}", "mean"),
        (r"\arg\min", "arg min"), (r"\left\|", "||"),
        (r"\right\|", "||"), (r"\{", "{"), (r"\}", "}"),
        (r"\ ", " "),
    ):
        value = value.replace(command, replacement)
    value = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", value)
    commands = {
        "ell": "ℓ", "Delta": "Δ", "Gamma": "Γ", "nabla": "∇",
        "partial": "∂", "approx": "≈", "lambda": "λ", "Sigma": "Σ",
        "sum": "Σ", "top": "T", "in": "∈", "to": "→", "le": "≤",
        "ge": "≥", "cdot": "·",
        "times": "×", "qquad": " ", "quad": " ", "left": "", "right": "",
        "operatorname": "", "mathrm": "", "text": "",
    }
    value = re.sub(r"\\([A-Za-z]+)", lambda m: commands.get(m.group(1), m.group(1)), value)
    value = re.sub(r"(_\{[^{}]*\}|_[A-Za-z0-9])(?=[A-Za-zΔΓΣ])", r"\1 ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value

def wrap_math_text(text: str, max_chars: int = 76) -> list[str]:
    """Wrap a display equation at top-level commas, then at arithmetic operators."""
    terms, start, stack = [], 0, []
    pairs = {")": "(", "]": "[", "}": "{"}
    for i, char in enumerate(text):
        if char in "([{":
            stack.append(char)
        elif char in ")]}":
            if stack and stack[-1] == pairs[char]:
                stack.pop()
        elif char in ",;" and not stack:
            terms.append(text[start:i].strip())
            start = i + 1
    terms.append(text[start:].strip())

    def split_long(term: str) -> list[str]:
        output = []
        while len(term) > max_chars:
            candidates = [term.rfind(op, 0, max_chars + 1) for op in (" + ", " - ", " * ", " / ")]
            position = max(candidates)
            if position >= max_chars // 2:
                operator = next(op for op in (" + ", " - ", " * ", " / ") if term.rfind(op, 0, max_chars + 1) == position)
                output.append(term[:position].rstrip() + operator.rstrip())
                term = term[position + len(operator):].lstrip()
                continue
            position = term.rfind(" ", 0, max_chars + 1)
            if position < 1:
                position = max_chars
                output.append(term[:position])
                term = term[position:]
            else:
                output.append(term[:position].rstrip())
                term = term[position:].lstrip()
        if term:
            output.append(term)
        return output

    lines, current = [], ""
    for term in terms:
        candidate = term if not current else current + ", " + term
        if len(candidate) <= max_chars:
            current = candidate
            continue
        if current:
            lines.append(current + ",")
            current = ""
        pieces = split_long(term)
        lines.extend(pieces[:-1])
        current = pieces[-1] if pieces else ""
    if current:
        lines.append(current)
    return lines or [text]

def math_run(text: str):
    run = OxmlElement("m:r")
    text_node = OxmlElement("m:t")
    text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    return run


def _math_group(text: str, start: int):
    opening = text[start]
    closing = {"{": "}", "(": ")", "[": "]"}[opening]
    depth = 0
    for index in range(start, len(text)):
        if text[index] == opening:
            depth += 1
        elif text[index] == closing:
            depth -= 1
            if depth == 0:
                return text[start + 1:index], index + 1
    return None, start


def _math_word_char(char: str) -> bool:
    import unicodedata
    return unicodedata.category(char)[0] in "LMN"


def _math_atom(text: str, start: int):
    if text.startswith("||", start):
        end = text.find("||", start + 2)
        if end >= 0:
            return text[start:end + 2], end + 2
    char = text[start]
    if char in "{([":
        inner, end = _math_group(text, start)
        if end > start:
            return text[start:end], end
    if _math_word_char(char):
        end = start + 1
        while end < len(text) and _math_word_char(text[end]):
            end += 1
        return text[start:end], end
    if text.startswith("->", start):
        return "->", start + 2
    return char, start + 1


def _append_math_atom(parent, atom: str):
    if atom.startswith("||") and atom.endswith("||") and len(atom) >= 4:
        parent.append(math_run("||"))
        append_math_expression(parent, atom[2:-2])
        parent.append(math_run("||"))
    elif atom[:1] in "{([" and atom[-1:] in "})]":
        parent.append(math_run(atom[0]))
        append_math_expression(parent, atom[1:-1])
        parent.append(math_run(atom[-1]))
    else:
        parent.append(math_run(atom))


def _script_content(text: str, start: int):
    if start >= len(text):
        return None, start
    if text[start] == "{":
        return _math_group(text, start)
    atom, end = _math_atom(text, start)
    return atom, end


def math_script(base: str, sub: str | None = None, sup: str | None = None):
    tag = "m:sSubSup" if sub is not None and sup is not None else "m:sSub" if sub is not None else "m:sSup"
    node = OxmlElement(tag)
    expr = OxmlElement("m:e")
    append_math_expression(expr, base)
    node.append(expr)
    if sub is not None:
        sub_node = OxmlElement("m:sub")
        append_math_expression(sub_node, sub)
        node.append(sub_node)
    if sup is not None:
        sup_node = OxmlElement("m:sup")
        append_math_expression(sup_node, sup)
        node.append(sup_node)
    return node


def append_math_expression(parent, text: str):
    """Build Word OMML script nodes for every subscript and superscript."""
    index = 0
    while index < len(text):
        if text[index].isspace():
            end = index + 1
            while end < len(text) and text[end].isspace():
                end += 1
            parent.append(math_run(text[index:end]))
            index = end
            continue
        atom, end = _math_atom(text, index)
        sub = sup = None
        while end < len(text) and text[end] in "_^":
            marker = text[end]
            content, after = _script_content(text, end + 1)
            if content is None or after <= end + 1:
                parent.append(math_run(atom + marker))
                atom = ""
                end += 1
                continue
            if marker == "_":
                sub = content
            else:
                sup = content
            end = after
        if atom:
            if sub is not None or sup is not None:
                parent.append(math_script(atom, sub=sub, sup=sup))
            else:
                _append_math_atom(parent, atom)
        index = end


def add_equation(doc, source: str, equation_number: str):
    text = math_text(source)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = 1  # centered on the text block
    table.autofit = False
    section = doc.sections[0]
    available_width = section.page_width - section.left_margin - section.right_margin
    side_width = Inches(0.36)
    widths = (side_width, available_width - side_width * 2, side_width)
    table_properties = table._tbl.tblPr
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:w"), str(int(available_width / 635)))
    table_width.set(qn("w:type"), "dxa")
    table_properties.append(table_width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    table_properties.append(borders)
    for index, width in enumerate(widths):
        table.columns[index].width = width
        cell = table.cell(0, index)
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_properties = cell._tc.get_or_add_tcPr()
        margins = OxmlElement("w:tcMar")
        for side in ("top", "left", "bottom", "right"):
            margin = OxmlElement(f"w:{side}")
            margin.set(qn("w:w"), "0")
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        cell_properties.append(margins)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_together = True

    math_cell = table.cell(0, 1)
    for line_index, line in enumerate(wrap_math_text(text)):
        math_paragraph = math_cell.paragraphs[0] if line_index == 0 else math_cell.add_paragraph()
        math_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math_paragraph.paragraph_format.space_before = Pt(0)
        math_paragraph.paragraph_format.space_after = Pt(0)
        math_paragraph.paragraph_format.line_spacing = 1.0
        math_paragraph.paragraph_format.keep_together = True
        office_math = OxmlElement("m:oMath")
        append_math_expression(office_math, line)
        math_paragraph._p.append(office_math)

    number_paragraph = table.cell(0, 2).paragraphs[0]
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_run = number_paragraph.add_run(f"({equation_number})")
    number_run.font.name = "Times New Roman"
    number_run.font.size = Pt(9)
    number_fonts = number_run._element.get_or_add_rPr().get_or_add_rFonts()
    number_fonts.set(qn("w:ascii"), "Times New Roman")
    number_fonts.set(qn("w:hAnsi"), "Times New Roman")
    row_properties = table.rows[0]._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:cantSplit"))


def shade_cell(cell, fill: str):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_cell_margins(cell, top=55, start=65, bottom=55, end=65):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def add_markdown_table(doc, lines):
    data = []
    for line in lines:
        cells = [x.strip() for x in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells):
            continue
        data.append(cells)
    if not data:
        return
    count = max(len(row) for row in data)
    table = doc.add_table(rows=len(data), cols=count)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(data):
        for col_index in range(count):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = row[col_index] if col_index < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if row_index == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell, "DCEBF2")
            elif col_index > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_font(run, size=7.7 if count >= 6 else 8.1, bold=(row_index == 0), color="17232D" if row_index == 0 else "202A32")
        if row_index == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_image(doc, src: str, image_root: Path):
    local = (image_root / src).resolve()
    if local.suffix.lower() == ".svg":
        local = local.with_suffix(".png")
    if not local.is_file():
        raise FileNotFoundError(local)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(local), width=Inches(6.75))


def add_reference(doc, index: int, fields):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(f"[{index}] {ieee_reference(fields)}")
    set_font(run, size=8.8, color="26343D")


def build(source: Path, output: Path):
    markdown = source.read_text(encoding="utf-8-sig")
    bib_block, entries = bibliography(markdown)
    body = markdown.replace(bib_block.group(0), "")
    order = cited_order(body, entries)
    fields_by_key = dict(entries)
    key_to_num = {key: i + 1 for i, key in enumerate(order)}

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    styles = doc.styles
    for name, size, color in [
        ("Normal", 10.3, "202A32"), ("Title", 18, "000000"),
        ("Heading 1", 14, "000000"), ("Heading 2", 11.5, "000000")
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    lines = body.splitlines()
    i = image_count = table_count = equation_count = 0
    references_inserted = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith(FENCE):
            while i < len(lines) and not lines[i].strip().startswith(FENCE):
                i += 1
            i += 1
            continue
        if line == "$$":
            i += 1
            expr = []
            while i < len(lines) and lines[i].strip() != "$$":
                expr.append(lines[i].strip())
                i += 1
            if i >= len(lines):
                raise ValueError("Unclosed equation block")
            equation_count += 1
            equation_source = " ".join(expr)
            tag = re.search(r"\\tag\{(\d+)\}", equation_source)
            equation_number = tag.group(1) if tag else str(equation_count)
            equation_source = re.sub(r"\\tag\{\d+\}", "", equation_source).strip()
            add_equation(doc, equation_source, equation_number)
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, table_lines)
            table_count += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level, title = len(heading.group(1)), heading.group(2).strip()
            if title == "References":
                doc.add_heading("References", level=1)
                for number, key in enumerate(order, 1):
                    add_reference(doc, number, fields_by_key[key])
                references_inserted = True
            elif level == 1:
                p = doc.add_paragraph(style="Title")
                run = p.add_run(title)
                set_font(run, size=18, bold=True, color="000000")
                p.paragraph_format.space_after = Pt(8)
            else:
                p = doc.add_paragraph(style="Heading 1" if level == 2 else "Heading 2")
                run = p.add_run(title)
                set_font(run, size=14 if level == 2 else 11.5, bold=True, color="000000")
                p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if line.startswith("![") and "](" in line and line.endswith(")"):
            source_path = line[2:].split("](", 1)[1][:-1]
            add_image(doc, source_path, source.parent)
            image_count += 1
            i += 1
            continue
        if line.startswith("*") and line.endswith("*") and line.count("*") >= 2:
            caption_text = line.strip("*")
            is_caption = caption_text.startswith(("Fig.", "Figure", "Table"))
            add_inline_paragraph(doc, caption_text, key_to_num, caption=is_caption, center=is_caption)
            i += 1
            continue
        if line.startswith("**") and line.endswith("**"):
            text = line[2:-2]
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_font(run, size=9.3, bold=True, color="202A32")
            i += 1
            continue
        if line.startswith("---"):
            i += 1
            continue

        paragraph = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or re.match(r"^#{1,3}\s", nxt) or nxt.startswith((FENCE, "![", "|", "$$", "*", "**", "---"))):
                break
            paragraph.append(nxt)
            i += 1
        joined = " ".join(paragraph).replace("\\ ", "").strip()
        add_inline_paragraph(doc, joined, key_to_num)

    if not references_inserted:
        doc.add_heading("References", level=1)
        for number, key in enumerate(order, 1):
            add_reference(doc, number, fields_by_key[key])
    doc.core_properties.title = "Patch-Content Fungibility: Geometry, Functional Transmission, and Operator-Aware Token Compression"
    doc.core_properties.subject = "Evidence-checked manuscript with IEEE-numbered references"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Created {output}; references={len(order)}; figures={image_count}; tables={table_count}; equations={equation_count}")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
